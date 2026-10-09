"""Assemble docs/pitch.pptx: LCT template slides 7–11, then the rendered pitch images.

Needs python-pptx 1.0 (pip install python-pptx==1.0.2). It is not a backend dependency.

The template is not stored in the repository. Pass LCT_PITCH_TEMPLATE or keep the file
in the owner's materials folder. Personal fields stay «заполнит владелец» until the owner
fills them. The PDF is exported from this file in PowerPoint, not by this script.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import sys
import zipfile
from copy import deepcopy
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
IMAGES = DOCS / "assets" / "pitch-slides"
OWNER = "заполнит владелец"

ESSENCE = "Работодатель сам находит человека и зовёт его с вилкой. Грейд подтверждает только тест."
UNIQUE = "Тест каждый раз с новыми числами. ФСП не меняет грейд. У каждой цифры есть команда проверки."
WHY = "Резюме не подтверждает навык, а протокол ФСП подтверждает. Обратный найм убирает ленту случайных откликов."
HARD = [
    "Одинаковые задания расходятся, поэтому числа в задаче каждый раз новые.",
    "Разряд — факт реестра, а не расчёт по месту. Цифры печатает команда.",
]
TECH = [
    "Адаптивный тест на IRT 3PL останавливается, когда грейд ясен.",
    "Задания параметрические: числа новые, сложность общая.",
    "Подбор раскладывается на навыки, грейд, тест и ось ФСП.",
    "Вход через ФСП ID: OpenID Connect с PKCE.",
    "Строка протокола — подписанное удостоверение.",
]
MARKET = [
    "Работодатель зовёт подходящих, а не читает ленту откликов.",
    "Вилка известна до начала общения.",
    "Спортивный результат виден как подтверждение, не как грейд.",
    "Нет истории ФСП — человек не теряет место.",
    "Дальше пилот с Федерацией: настоящий ФСП ID и реестр.",
]


def template_path() -> Path:
    env = os.environ.get("LCT_PITCH_TEMPLATE")
    if env:
        path = Path(env)
        if not path.is_file():
            raise SystemExit("LCT_PITCH_TEMPLATE не указывает на файл шаблона")
        return path
    desktop = Path.home() / "Desktop" / "Талант" / "ЛЦТ_2026_Шаблон презентации ФСП.pptx"
    if desktop.is_file():
        return desktop
    raise SystemExit("Нет шаблона ЛЦТ. Задайте переменную LCT_PITCH_TEMPLATE.")


def delete_slide(prs: Presentation, index: int) -> None:
    slide_id = prs.slides._sldIdLst[index]
    prs.part.drop_rel(slide_id.get(qn("r:id")))
    prs.slides._sldIdLst.remove(slide_id)


def shape_starting(slide, prefix: str):
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip().startswith(prefix):
            return shape
    raise SystemExit(f"на слайде нет текста «{prefix}»")


def set_run(paragraph, text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.text = text
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def set_value(paragraph, value: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.text = value
        return
    if len(runs) >= 2:
        runs[-1].text = value
        for run in runs[1:-1]:
            run.text = ""
        return
    label = runs[0].text.rstrip()
    runs[0].text = f"{label} {value}" if label.endswith(":") else value


def fill_lines(paragraph, lines: list[str]) -> None:
    """Keep the first run's formatting and put one visual line per sentence."""
    proto = paragraph.runs[0]._r
    host = paragraph._p
    for child in list(host):
        if child.tag != qn("a:pPr"):
            host.remove(child)
    for index, line in enumerate(lines):
        if index:
            broken = etree.Element(qn("a:br"))
            props = proto.find(qn("a:rPr"))
            if props is not None:
                broken.append(deepcopy(props))
            host.append(broken)
        run = deepcopy(proto)
        node = run.find(qn("a:t"))
        node.text = line
        host.append(run)


def fill_bullets(shape, lines: list[str]) -> None:
    """One template paragraph per line, with the first paragraph's bullet properties."""
    frame = shape.text_frame
    frame.word_wrap = True
    body = frame._txBody
    first = body.find(qn("a:p"))
    if first is None:
        raise SystemExit("в поле слайда нет абзаца")
    bullet = first.find(qn("a:pPr"))
    proto = first.find(qn("a:r"))
    for paragraph in list(body.findall(qn("a:p"))):
        body.remove(paragraph)
    for line in lines:
        paragraph = etree.Element(qn("a:p"))
        if bullet is not None:
            paragraph.append(deepcopy(bullet))
        run = deepcopy(proto) if proto is not None else etree.Element(qn("a:r"))
        node = run.find(qn("a:t"))
        if node is None:
            node = etree.SubElement(run, qn("a:t"))
        node.text = line
        paragraph.append(run)
        body.append(paragraph)


def fill_shape(shape, lines: list[str]) -> None:
    frame = shape.text_frame
    frame.word_wrap = True
    paragraphs = list(frame.paragraphs)
    fill_lines(paragraphs[0], lines)
    for paragraph in paragraphs[1:]:
        if paragraph.runs:
            set_run(paragraph, "")


def notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def fill_banner(slide, text: str) -> None:
    """The pink title pill is an empty placeholder on the team slides. The colour comes from the layout."""
    for shape in slide.shapes:
        if shape.has_text_frame and shape.name.startswith("Заголовок") and not shape.text_frame.text.strip():
            set_run(shape.text_frame.paragraphs[0], text)
            return
    raise SystemExit("нет пустого заголовка на слайде шаблона")


def fill_template(prs: Presentation) -> None:
    cover, team, people, story, short = list(prs.slides)[:5]
    title = None
    for shape in cover.shapes:
        if shape.has_text_frame and not shape.text_frame.text.strip() and shape.name.startswith("Заголовок"):
            title = shape
    if title is not None:
        set_run(title.text_frame.paragraphs[0], "ФСП Талант")
    notes(cover, "Название задачи уже на слайде. Название команды владелец впишет, если оно отличается от названия решения.")

    set_run(shape_starting(team, "В чем суть").text_frame.paragraphs[0], ESSENCE)
    set_run(shape_starting(team, "Что делает").text_frame.paragraphs[0], UNIQUE)
    about = shape_starting(team, "Капитан")
    paragraphs = list(about.text_frame.paragraphs)
    set_value(paragraphs[0], OWNER)
    set_value(paragraphs[1], OWNER)
    set_value(paragraphs[2], OWNER)
    set_run(paragraphs[3], "")
    set_run(paragraphs[4], "")
    set_value(paragraphs[5], OWNER)
    body = about.text_frame._txBody
    for paragraph in list(body.findall(qn("a:p"))):
        if not "".join(node.text or "" for node in paragraph.findall(".//" + qn("a:t"))).strip():
            body.remove(paragraph)
    notes(team, "Суть и уникальность говорим своими словами. Блок «О команде» владелец заполняет перед сдачей.")

    for shape in people.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip().startswith("Имя Фамилия"):
            set_run(shape.text_frame.paragraphs[0], OWNER)
    fill_banner(people, "Команда")
    notes(people, "Пять карточек без имён: роли на месте, имена и контакты присылает владелец.")
    fill_banner(story, "Как шла работа")

    set_run(shape_starting(story, "Расскажите, как вы собрались").text_frame.paragraphs[0], OWNER)
    set_run(shape_starting(story, "Что вас вдохновило").text_frame.paragraphs[0], WHY)
    fill_shape(shape_starting(story, "Расскажите о самых интересных"), HARD)
    notes(story, "Историю команды заполняет владелец. Почему эта задача и какие были сложности — уже вписаны.")

    fill_bullets(shape_starting(short, "Опишите в чем техническая"), TECH)
    fill_bullets(shape_starting(short, "Опишите ваши идеи"), MARKET)
    notes(short, "Техническая суть слева, пилот с Федерацией справа. На защите этот слайд — последние секунды.")


def add_images(prs: Presentation) -> None:
    manifest = json.loads((IMAGES / "manifest.json").read_text(encoding="utf-8"))
    blank = prs.slide_layouts[22]
    for item in manifest:
        slide = prs.slides.add_slide(blank)
        picture = slide.shapes.add_picture(
            str(IMAGES / item["file"]), Emu(0), Emu(0), prs.slide_width, prs.slide_height)
        picture._element.nvPicPr.cNvPr.set("descr", item["alt"])
        notes(slide, item["notes"])


def stamp(prs: Presentation) -> None:
    props = prs.core_properties
    props.author = "Konkov Dmitrij"
    props.last_modified_by = "Konkov Dmitrij"
    props.title = "ФСП Талант — обратный найм"
    props.subject = "ЛЦТ 2026, трек Федерации спортивного программирования"
    props.category = ""
    props.keywords = ""
    html = (DOCS / "pitch.html").read_bytes().replace(b"\r\n", b"\n")
    props.comments = "pitch-html-sha256=" + hashlib.sha256(html).hexdigest()


def patch_package(path: Path, slides: int) -> None:
    """python-pptx leaves the template's app.xml, which still says 37 slides and 4:3."""
    titles = "".join(f"<vt:lpstr>Slide {n}</vt:lpstr>" for n in range(1, slides + 1))
    with zipfile.ZipFile(path) as source:
        files = {name: source.read(name) for name in source.namelist()}
    app = files["docProps/app.xml"].decode("utf-8")
    app = app.replace("<Slides>37</Slides>", f"<Slides>{slides}</Slides>")
    app = app.replace("<Notes>37</Notes>", f"<Notes>{slides}</Notes>")
    app = app.replace("On-screen Show (4:3)", "On-screen Show (16:9)")
    app = app.replace("<vt:i4>37</vt:i4>", f"<vt:i4>{slides}</vt:i4>")
    app = re.sub(
        r"<vt:vector size=\"38\" baseType=\"lpstr\">.*?</vt:vector>",
        f'<vt:vector size="{slides + 1}" baseType="lpstr"><vt:lpstr>Theme 1</vt:lpstr>{titles}</vt:vector>',
        app,
        count=1,
    )
    files["docProps/app.xml"] = app.encode("utf-8")
    core = files["docProps/core.xml"].decode("utf-8")
    core = re.sub(
        r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
        rf"\g<1>{datetime.now(UTC).strftime('%Y-%m-%dT%H:%M:%SZ')}\g<2>",
        core,
        count=1,
    )
    files["docProps/core.xml"] = core.encode("utf-8")
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for name, blob in files.items():
            target.writestr(name, blob)
    path.write_bytes(buffer.getvalue())


def main() -> None:
    if not (IMAGES / "manifest.json").is_file():
        raise SystemExit("Нет кадров слайдов. Сначала: node docs/tools/build_pitch.mjs")
    prs = Presentation(str(template_path()))
    keep = {6, 7, 8, 9, 10}
    for index in range(len(prs.slides) - 1, -1, -1):
        if index not in keep:
            delete_slide(prs, index)
    if len(prs.slides) != 5:
        raise SystemExit(f"после очистки осталось {len(prs.slides)} слайдов, ждали 5")
    # New slides are named from the current count. Rename first, or slide7.xml collides
    # with a template slide that still has that part name.
    prs.part.rename_slide_parts([slide_id.get(qn("r:id")) for slide_id in prs.slides._sldIdLst])
    fill_template(prs)
    add_images(prs)
    stamp(prs)
    out = DOCS / "pitch.pptx"
    count = len(prs.slides)
    prs.save(str(out))
    patch_package(out, count)
    saved = Presentation(str(out))
    if len(saved.slides) != 16:
        raise SystemExit(f"в файле {len(saved.slides)} слайдов, ждали 16")
    joined = "\n".join(
        shape.text_frame.text
        for slide in saved.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    for stale in ("Имя Фамилия", "ФИО, специальность", "__ человек", "Опишите в чем", "Опишите ваши идеи", "как образовалась"):
        if stale in joined:
            raise SystemExit(f"в презентации осталась заготовка шаблона «{stale}»")
    if joined.count(OWNER) < 5:
        raise SystemExit("нет пометки «заполнит владелец»")
    print(f"pitch.pptx: {count} slides")


if __name__ == "__main__":
    sys.exit(main())
